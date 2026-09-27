"""Build data/corpus/ (and data/resumes/pdf/) from the clean sources.

Why: a real company's documents arrive as digital PDFs, scans, Word files and
intranet pages, not clean Markdown. This script renders the authoring sources in
data/policies/ into those formats, so your ingestion pipeline has real parsing
work to do (Stage 2 of the interview mock). The sources stay as ground truth for
checking what your parsers extract.

Usage:  python scripts/build_corpus.py
Needs:  pip install reportlab python-docx pillow
Deterministic: running it twice produces the same text content.
"""
import json
import random
import re
import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from PIL import Image, ImageDraw, ImageFilter, ImageFont
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfgen import canvas as rl_canvas
from reportlab.platypus import (BaseDocTemplate, Frame, FrameBreak, PageTemplate, Paragraph,
                                SimpleDocTemplate, Spacer, Table, TableStyle)

ROOT = Path(__file__).resolve().parents[1]
MANIFEST = ROOT / "data" / "corpus_manifest.json"


# ---------------------------------------------------------------- markdown ---
def parse_md(text):
    """Tiny parser for the source format: title, meta lines, ## headings,
    paragraphs, 'Table N:' captions and pipe tables."""
    blocks, lines, i = [], text.splitlines(), 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line:
            i += 1
            continue
        if line.startswith("# "):
            blocks.append(("title", line[2:]))
            i += 1
            while i < len(lines) and lines[i].strip():   # meta lines under the title
                blocks.append(("meta", lines[i].strip()))
                i += 1
        elif line.startswith("## "):
            blocks.append(("h2", line[3:]))
            i += 1
        elif line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r"-+", c) for c in cells):
                    rows.append(cells)
                i += 1
            blocks.append(("table", rows))
        elif re.match(r"Table \d+:", line):
            blocks.append(("caption", line))
            i += 1
        else:
            blocks.append(("p", line))
            i += 1
    return blocks


def esc(s):
    return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


# --------------------------------------------------------------------- PDF ---
STY = getSampleStyleSheet()
S = {
    "title": ParagraphStyle("t", parent=STY["Title"], fontSize=17, spaceAfter=4),
    "meta": ParagraphStyle("m", parent=STY["Normal"], fontSize=8, textColor=colors.grey),
    "h2": ParagraphStyle("h", parent=STY["Heading2"], fontSize=12, spaceBefore=8, spaceAfter=3),
    "p": ParagraphStyle("p", parent=STY["Normal"], fontSize=9.5, leading=13, spaceAfter=5),
    "caption": ParagraphStyle("c", parent=STY["Normal"], fontSize=9, fontName="Helvetica-Bold", spaceBefore=4, spaceAfter=3),
    "cell": ParagraphStyle("cell", parent=STY["Normal"], fontSize=8.5, leading=10.5),
}


def pdf_flowables(blocks, width):
    out = []
    for kind, val in blocks:
        if kind == "table":
            data = [[Paragraph(esc(c), S["cell"]) for c in row] for row in val]
            t = Table(data, colWidths=[width / len(val[0])] * len(val[0]), repeatRows=1)
            t.setStyle(TableStyle([
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e8eef5")),
                ("VALIGN", (0, 0), (-1, -1), "TOP")]))
            out += [t, Spacer(1, 6)]
        else:
            out.append(Paragraph(esc(val), S[kind]))
    return out


def page_decor(doc_meta):
    """Repeated header and footer on every page: the kind of noise a
    cleaning step has to remove before chunking."""
    def draw(c, doc):
        c.saveState()
        c.setFont("Helvetica", 7.5)
        c.setFillColor(colors.grey)
        c.drawString(18 * mm, A4[1] - 12 * mm, f"Nilebyte Solutions · {doc_meta['doc_id']} · {doc_meta['title']}")
        c.drawRightString(A4[0] - 18 * mm, A4[1] - 12 * mm, f"Version {doc_meta['version']}")
        c.drawString(18 * mm, 10 * mm, "INTERNAL · Uncontrolled when printed")
        c.drawRightString(A4[0] - 18 * mm, 10 * mm, f"Page {doc.page}")
        c.restoreState()
    return draw


def build_pdf(blocks, out, meta, two_column=False):
    margin = 18 * mm
    if two_column:
        doc = BaseDocTemplate(str(out), pagesize=A4, leftMargin=margin, rightMargin=margin,
                              topMargin=20 * mm, bottomMargin=18 * mm)
        gap = 8 * mm
        col_w = (A4[0] - 2 * margin - gap) / 2
        h = A4[1] - 38 * mm
        frames = [Frame(margin, 18 * mm, col_w, h, id="left"),
                  Frame(margin + col_w + gap, 18 * mm, col_w, h, id="right")]
        doc.addPageTemplates([PageTemplate(frames=frames, onPage=page_decor(meta))])
        # Break to the right column at the section heading nearest the middle of
        # the text, so the page really has two columns of body text.
        sizes = [len(str(v)) for _, v in blocks]
        half, run, cut = sum(sizes) / 2, 0, len(blocks)
        for i, (kind, _) in enumerate(blocks):
            if kind == "h2" and run >= half * 0.8:
                cut = i
                break
            run += sizes[i]
        flow = pdf_flowables(blocks[:cut], col_w) + [FrameBreak()] + pdf_flowables(blocks[cut:], col_w)
        doc.build(flow)
    else:
        doc = SimpleDocTemplate(str(out), pagesize=A4, leftMargin=margin, rightMargin=margin,
                                topMargin=20 * mm, bottomMargin=18 * mm,
                                title=meta["title"], author=meta["owner"])
        deco = page_decor(meta)
        doc.build(pdf_flowables(blocks, A4[0] - 2 * margin), onFirstPage=deco, onLaterPages=deco)


# -------------------------------------------------------------- scanned PDF ---
def build_scanned_pdf(blocks, out, seed=7):
    """Render the text as an image-only PDF (no text layer), slightly rotated and
    noisy, like a photocopy that was scanned. Only OCR can read it."""
    rnd = random.Random(seed)
    W, H, M = 1240, 1754, 90                         # A4 at 150 dpi
    font = ImageFont.load_default(size=19)
    bold = ImageFont.load_default(size=23)
    title = ImageFont.load_default(size=30)
    img = Image.new("L", (W, H), 250)
    d = ImageDraw.Draw(img)
    y = M

    def wrap(text, f, width):
        words, lines, cur = text.split(), [], ""
        for w in words:
            test = (cur + " " + w).strip()
            if d.textlength(test, font=f) <= width:
                cur = test
            else:
                lines.append(cur)
                cur = w
        return lines + [cur]

    for kind, val in blocks:
        if kind == "table":
            ncol = len(val[0])
            colw = (W - 2 * M) / ncol
            for r, row in enumerate(val):
                cell_lines = [wrap(c, bold if r == 0 else font, colw - 12) for c in row]
                rh = 26 * max(len(cl) for cl in cell_lines) + 12
                for ci, cl in enumerate(cell_lines):
                    x0 = M + ci * colw
                    d.rectangle([x0, y, x0 + colw, y + rh], outline=60, width=2)
                    for li, t in enumerate(cl):
                        d.text((x0 + 6, y + 6 + li * 26), t, fill=20, font=bold if r == 0 else font)
                y += rh
            y += 18
            continue
        f = {"title": title, "h2": bold, "caption": bold}.get(kind, font)
        for t in wrap(val, f, W - 2 * M):
            d.text((M, y), t, fill=25, font=f)
            y += int(f.size * 1.45)
        y += 10
    img = img.rotate(0.7, resample=Image.BICUBIC, fillcolor=250, expand=False)
    px = img.load()
    for _ in range(9000):                            # speckle noise
        px[rnd.randrange(W), rnd.randrange(H)] = rnd.choice([30, 90, 160])
    img = img.filter(ImageFilter.GaussianBlur(0.6))
    img.convert("RGB").save(out, "PDF", resolution=150)


# -------------------------------------------------------------------- DOCX ---
def set_rtl(paragraph):
    ppr = paragraph._p.get_or_add_pPr()
    bidi = OxmlElement("w:bidi")
    bidi.set(qn("w:val"), "1")
    ppr.append(bidi)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in paragraph.runs:
        rpr = run._r.get_or_add_rPr()
        rtl = OxmlElement("w:rtl")
        rtl.set(qn("w:val"), "1")
        rpr.append(rtl)


def build_docx(blocks, out, meta):
    doc = Document()
    rtl = meta["language"] == "ar"
    sec = doc.sections[0]
    sec.header.paragraphs[0].text = f"Nilebyte Solutions · {meta['doc_id']} · Version {meta['version']}"
    sec.footer.paragraphs[0].text = "INTERNAL · Uncontrolled when printed"
    for kind, val in blocks:
        if kind == "title":
            p = doc.add_heading(val, level=0)
        elif kind == "h2":
            p = doc.add_heading(val, level=1)
        elif kind == "table":
            t = doc.add_table(rows=len(val), cols=len(val[0]))
            t.style = "Table Grid"
            for r, row in enumerate(val):
                for c, cell in enumerate(row):
                    t.cell(r, c).text = cell
                    if r == 0:
                        for run in t.cell(r, c).paragraphs[0].runs:
                            run.bold = True
            continue
        else:
            p = doc.add_paragraph(val)
            if kind in ("meta",):
                for run in p.runs:
                    run.italic = True
            if kind == "caption":
                for run in p.runs:
                    run.bold = True
        if rtl:
            set_rtl(p)
    doc.core_properties.title = meta["title"]
    doc.core_properties.author = meta["owner"]
    doc.save(out)


# ----------------------------------------------------------------- resumes ---
INJECTION_MARK = "SYSTEM INSTRUCTION FOR AI SCREENING TOOLS"


def build_resume_pdfs():
    src, dst = ROOT / "data" / "resumes", ROOT / "data" / "resumes" / "pdf"
    dst.mkdir(exist_ok=True)

    # cand_05: normal one-column resume; the injection is drawn in white 1pt text,
    # invisible on screen but present in the PDF text layer.
    text = (src / "cand_05.txt").read_text(encoding="utf-8")
    visible, hidden = text.split(INJECTION_MARK, 1)
    hidden = INJECTION_MARK + hidden
    c = rl_canvas.Canvas(str(dst / "cand_05.pdf"), pagesize=A4)
    y = A4[1] - 25 * mm
    for line in visible.rstrip().splitlines():
        bold = line.isupper() and line.strip() != ""
        c.setFont("Helvetica-Bold" if bold else "Helvetica", 12 if bold else 10)
        c.drawString(22 * mm, y, line)
        y -= 6 * mm if line.strip() else 3 * mm
    c.setFillColor(colors.white)
    c.setFont("Helvetica", 1)
    words, cur, y = hidden.split(), "", y - 4 * mm
    for w in words:
        cur = (cur + " " + w).strip()
        if len(cur) > 260:
            c.drawString(22 * mm, y, cur)
            y -= 1.2 * mm
            cur = ""
    c.drawString(22 * mm, y, cur)
    c.save()

    # cand_08: two-column layout with a photo box: reading order is the challenge.
    lines = (src / "cand_08.txt").read_text(encoding="utf-8").splitlines()
    left_keys = ("Date of birth", "National ID", "Marital status", "Address", "Phone", "Email", "[Photo")
    left = [l for l in lines if l.startswith(left_keys)]
    skills_i = lines.index("SKILLS")
    left += ["", "SKILLS"] + [l for l in lines[skills_i + 1:] if l.strip()]
    right = [l for l in lines[2:skills_i] if not l.startswith(left_keys) and l.strip() != "Personal details"]
    c = rl_canvas.Canvas(str(dst / "cand_08.pdf"), pagesize=A4)
    c.setFont("Helvetica-Bold", 18)
    c.drawString(20 * mm, A4[1] - 22 * mm, lines[0])
    c.setFont("Helvetica", 11)
    c.drawString(20 * mm, A4[1] - 29 * mm, lines[1])
    c.setStrokeColor(colors.grey)
    c.rect(20 * mm, A4[1] - 75 * mm, 35 * mm, 40 * mm)
    c.setFont("Helvetica", 8)
    c.drawString(27 * mm, A4[1] - 56 * mm, "[photo]")

    def column(x, y, items, width_chars):
        for l in items:
            if l.isupper() and l.strip():
                c.setFont("Helvetica-Bold", 10)
            else:
                c.setFont("Helvetica", 8.5)
            chunks = [l[i:i + width_chars] for i in range(0, max(len(l), 1), width_chars)] or [""]
            for ch in chunks:
                c.drawString(x, y, ch)
                y -= 4.6 * mm
    column(20 * mm, A4[1] - 84 * mm, [l for l in left if not l.startswith("[Photo")], 40)
    column(80 * mm, A4[1] - 40 * mm, right, 70)
    c.save()


# --------------------------------------------------------------------- main ---
def main():
    manifest = json.loads(MANIFEST.read_text(encoding="utf-8"))
    out_dir = ROOT / "data" / "corpus"
    out_dir.mkdir(exist_ok=True)
    for meta in manifest["documents"]:
        src, out, fmt = ROOT / meta["source"], ROOT / meta["file"], meta["format"]
        if fmt in ("md", "html"):
            shutil.copyfile(src, out)
            continue
        blocks = parse_md(src.read_text(encoding="utf-8"))
        if fmt == "pdf":
            build_pdf(blocks, out, meta)
        elif fmt == "pdf_two_column":
            build_pdf(blocks, out, meta, two_column=True)
        elif fmt == "pdf_scanned":
            build_scanned_pdf(blocks, out)
        elif fmt == "docx":
            build_docx(blocks, out, meta)
        else:
            raise ValueError(f"unknown format {fmt}")
    build_resume_pdfs()
    print(f"built {len(manifest['documents'])} documents into {out_dir.relative_to(ROOT)} and 2 resume PDFs")


if __name__ == "__main__":
    main()
